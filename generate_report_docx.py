import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    PRIMARY_COLOR = RGBColor(79, 70, 229)     # #4F46E5 (Indigo)
    SECONDARY_COLOR = RGBColor(30, 41, 59)   # #1E293B (Slate 800)
    MUTED_COLOR = RGBColor(100, 116, 139)    # #64748B
    ACCENT_COLOR = RGBColor(16, 185, 129)    # #10B981 (Emerald)

    # Header title
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("FPT UNIVERSITY – DEPARTMENT OF COMPUTING\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = MUTED_COLOR

    r_course = p_inst.add_run("MMA301 – MOBILE MULTI-PLATFORM APPLICATION DEVELOPMENT\n")
    r_course.font.name = "Arial"
    r_course.font.size = Pt(14)
    r_course.font.bold = True
    r_course.font.color.rgb = SECONDARY_COLOR

    r_title = p_inst.add_run("PRACTICAL EXAM 2 REPORT\nAUTHENTICATION, TEAMS & REAL-TIME CHAT")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    # Horizontal divider
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("━" * 55)
    r_div.font.color.rgb = PRIMARY_COLOR

    # --- SECTION 1: STUDENT & PROJECT INFORMATION ---
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("I. THÔNG TIN SINH VIÊN & ĐƯỜNG DẪN DỰ ÁN")
    h1_run.font.name = "Arial"
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = PRIMARY_COLOR

    # Student Info Table
    table_info = doc.add_table(rows=5, cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_info.autofit = False

    info_data = [
        ("Họ và tên sinh viên:", "[ĐIỀN TÊN CỦA BẠN - VÍ DỤ: NGUYỄN VĂN QUÝ]"),
        ("Mã số sinh viên (StudentID):", "[ĐIỀN MSSV CỦA BẠN - VÍ DỤ: QE170123 / SE160999]"),
        ("Lớp học (Class):", "[ĐIỀN LỚP - VÍ DỤ: MMA301.SE1701]"),
        ("Giảng viên hướng dẫn:", "[ĐIỀN TÊN GIẢNG VIÊN CỦA BẠN]"),
        ("Ngày nộp bài:", "16/10/2026")
    ]

    for i, (label, val) in enumerate(info_data):
        row = table_info.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = SECONDARY_COLOR

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        if "[" in val:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(220, 38, 38) # Red reminder
        else:
            r1.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph()

    # Project Links Table
    p_link_hdr = doc.add_paragraph()
    r_lh = p_link_hdr.add_run("🔗 Các liên kết dự án (Project Links):")
    r_lh.font.name = "Arial"
    r_lh.font.bold = True
    r_lh.font.size = Pt(11)

    table_links = doc.add_table(rows=3, cols=2)
    table_links.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_links.autofit = False

    links_data = [
        ("Public GitHub Repository:", "https://github.com/DevTurtle-Hub/ASSIGNMENT_MMA301"),
        ("EAS Build APK Download Link:", "https://expo.dev/accounts/quynv64/projects/task-management-app/builds/4735525b-ae79-4fcb-a46e-52221f2f2921"),
        ("Expo Project Link / QR Code:", "exp://u.expo.dev/e3458b13-6cf6-4bda-9bd8-d3c012522a39 (hoặc link Expo web)")
    ]

    for i, (label, val) in enumerate(links_data):
        row = table_links.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        
        set_cell_background(c0, "EEF2FF")
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = PRIMARY_COLOR

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph()

    # --- SECTION 2: ARCHITECTURE & FIRESTORE DATABASE SCHEMA ---
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("II. KIẾN TRÚC HỆ THỐNG & FIRESTORE DATA SCHEMA (ERD)")
    h2_run.font.name = "Arial"
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = PRIMARY_COLOR

    p_schema_intro = doc.add_paragraph()
    p_schema_intro.add_run(
        "Hệ thống ứng dụng Task Management – Practical Exam 2 được phát triển trên nền tảng React Native (Expo) và Firebase (Authentication & Cloud Firestore). Cơ sở dữ liệu được tổ chức theo mô hình quan hệ đa thực thể phân tán:"
    )

    schema_tables = [
        ("1. Collection: `users` (Thông tin tài khoản người dùng)", [
            ("id", "string (Primary Key)", "Khóa chính, trùng với UID từ Firebase Authentication."),
            ("name", "string", "Họ và tên đầy đủ của người dùng."),
            ("email", "string", "Địa chỉ email định danh đăng nhập."),
            ("avatarUrl", "string", "Đường dẫn ảnh đại diện (avatar)."),
            ("createdAt", "string (ISO 8601)", "Thời điểm tài khoản được khởi tạo.")
        ]),
        ("2. Collection: `teams` (Thông tin nhóm làm việc)", [
            ("id", "string (Primary Key)", "Mã định danh nhóm sinh tự động bởi Firestore."),
            ("name", "string", "Tên nhóm làm việc (ví dụ: Mobile Dev Squad)."),
            ("description", "string", "Mô tả mục tiêu của nhóm."),
            ("ownerId", "string (Foreign Key)", "Khóa ngoại liên kết tới `users.id` của người tạo nhóm."),
            ("code", "string (Unique)", "Mã Code 6 ký tự để thành viên khác nhập và tham gia nhóm."),
            ("createdAt", "string (ISO 8601)", "Thời điểm tạo nhóm.")
        ]),
        ("3. Collection: `teamMembers` (Quan hệ thành viên trong nhóm)", [
            ("id", "string (Primary Key)", "Mã định danh quan hệ thành viên."),
            ("teamId", "string (Foreign Key)", "Khóa ngoại tham chiếu tới `teams.id`."),
            ("userId", "string (Foreign Key)", "Khóa ngoại tham chiếu tới `users.id`."),
            ("userName", "string", "Tên hiển thị của thành viên."),
            ("userEmail", "string", "Email của thành viên."),
            ("role", "string", "Vai trò trong nhóm (`owner` hoặc `member`)."),
            ("joinedAt", "string (ISO 8601)", "Thời điểm gia nhập nhóm.")
        ]),
        ("4. Collection: `tasks` (Công việc được liên kết Team & Assignee)", [
            ("id", "string (Primary Key)", "Mã định danh công việc (Firestore Auto-ID)."),
            ("title", "string", "Tiêu đề công việc cần làm."),
            ("description", "string", "Chi tiết nội dung công việc."),
            ("status", "string", "Trạng thái tiến độ: `todo`, `in_progress`, `done`."),
            ("priority", "string", "Mức độ ưu tiên: `low`, `medium`, `high`."),
            ("dueDate", "string", "Hạn hoàn thành (YYYY-MM-DD)."),
            ("createdAt", "string (ISO 8601)", "Thời điểm tạo task."),
            ("teamId", "string | null", "Khóa ngoại liên kết nhóm (nếu gán cho team)."),
            ("assigneeId", "string | null", "Khóa ngoại liên kết thành viên được giao task."),
            ("teamName", "string | null", "Tên nhóm được gán (denormalized for display)."),
            ("assigneeName", "string | null", "Tên người được gán (denormalized for display).")
        ]),
        ("5. Subcollection: `teams/{teamId}/messages` (Tin nhắn trao đổi nội bộ nhóm)", [
            ("id", "string (Primary Key)", "Mã tin nhắn tự sinh trong subcollection."),
            ("senderId", "string (Foreign Key)", "Khóa ngoại tham chiếu tới `users.id` của người gửi."),
            ("senderName", "string", "Tên của người gửi tin nhắn."),
            ("text", "string", "Nội dung văn bản tin nhắn chat."),
            ("createdAt", "string (ISO 8601)", "Thời điểm gửi tin nhắn (dùng để sắp xếp real-time).")
        ])
    ]

    for title, fields in schema_tables:
        p_t = doc.add_paragraph()
        r_t = p_t.add_run(title)
        r_t.font.name = "Arial"
        r_t.font.bold = True
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = SECONDARY_COLOR

        tbl = doc.add_table(rows=len(fields)+1, cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Header row
        hdr = tbl.rows[0]
        hdr.cells[0].width = Inches(1.8)
        hdr.cells[1].width = Inches(2.2)
        hdr.cells[2].width = Inches(3.0)
        for idx, text in enumerate(["Tên trường (Field)", "Kiểu dữ liệu (Type)", "Mô tả / Ràng buộc"]):
            cell = hdr.cells[idx]
            set_cell_background(cell, "334155")
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, (f_name, f_type, f_desc) in enumerate(fields):
            row = tbl.rows[r_idx+1]
            row.cells[0].width = Inches(1.8)
            row.cells[1].width = Inches(2.2)
            row.cells[2].width = Inches(3.0)

            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            for c_idx, val in enumerate([f_name, f_type, f_desc]):
                cell = row.cells[c_idx]
                set_cell_background(cell, bg)
                set_cell_margins(cell, 50, 50, 80, 80)
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(9)
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = PRIMARY_COLOR

        doc.add_paragraph()

    # Firestore Security Rules section
    p_sec = doc.add_paragraph()
    r_sec = p_sec.add_run("Quy tắc bảo mật Firestore (Firestore Security Rules):")
    r_sec.font.name = "Arial"
    r_sec.font.bold = True
    r_sec.font.size = Pt(11)

    p_sec_code = doc.add_paragraph()
    p_sec_code.paragraph_format.left_indent = Inches(0.2)
    p_sec_code.paragraph_format.right_indent = Inches(0.2)
    r_code = p_sec_code.add_run(
        "rules_version = '2';\n"
        "service cloud.firestore {\n"
        "  match /databases/{database}/documents {\n"
        "    function isAuthenticated() {\n"
        "      return request.auth != null;\n"
        "    }\n"
        "    match /users/{userId} { allow read: if isAuthenticated(); allow write: if isAuthenticated() && request.auth.uid == userId; }\n"
        "    match /teams/{teamId} {\n"
        "      allow read, write: if isAuthenticated();\n"
        "      match /messages/{messageId} { allow read, write: if isAuthenticated(); }\n"
        "    }\n"
        "    match /teamMembers/{memberId} { allow read, write: if isAuthenticated(); }\n"
        "    match /tasks/{taskId} { allow read, write: if isAuthenticated(); }\n"
        "  }\n"
        "}"
    )
    r_code.font.name = "Courier New"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph()

    # --- SECTION 3: 11 SCREENSHOTS AND CAPTIONS ---
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("III. DANH SÁCH 11 ẢNH MINH CHỨNG & CHÚ THÍCH (CAPTIONS)")
    h3_run.font.name = "Arial"
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = PRIMARY_COLOR

    p_sc_guide = doc.add_paragraph()
    p_sc_guide.add_run(
        "Dưới đây là 11 ảnh minh chứng chi tiết theo đúng quy định của đề bài Practical Exam 2. Mỗi ảnh đã được thiết kế sẵn khung giữ chỗ (placeholder), hướng dẫn thao tác chụp và dòng chú thích (Caption) quy chuẩn:"
    )

    screenshots_info = [
        (
            1,
            "Ảnh 1: Firebase Authentication Console",
            "Mở trình duyệt vào Firebase Console -> Authentication -> tab Users. Hiển thị danh sách các tài khoản test đã tạo (ít nhất 2 tài khoản, ví dụ: alice@gmail.com và bob@gmail.com).",
            "Bảng điều khiển Firebase Authentication hiển thị danh sách các tài khoản test đã đăng ký (alice@gmail.com, bob@gmail.com).",
            "Đáp ứng yêu cầu 1 (Firebase Authentication): Tài khoản người dùng được lưu trữ an toàn trong dịch vụ Authentication của Firebase, hỗ trợ đăng nhập Email/Password."
        ),
        (
            2,
            "Ảnh 2: Firestore Database Console",
            "Mở Firebase Console -> Firestore Database -> tab Data. Chụp toàn cảnh hiển thị các root collections: users, teams, teamMembers, tasks và subcollection teams/{teamId}/messages kèm dữ liệu mẫu.",
            "Bảng điều khiển Cloud Firestore hiển thị đầy đủ các collections users, teams, teamMembers, tasks và subcollection messages cùng dữ liệu thực tế.",
            "Đáp ứng yêu cầu 1, 2, 3, 4: Cơ sở dữ liệu NoSQL Cloud Firestore đồng bộ real-time, lưu trữ toàn bộ thực thể của hệ thống."
        ),
        (
            3,
            "Ảnh 3: Firestore Security Rules",
            "Mở Firebase Console -> Firestore Database -> tab Rules. Chụp toàn màn hình hiển thị nội dung file firestore.rules đã được Publish thành công.",
            "Quy tắc bảo mật Firestore Security Rules yêu cầu request.auth != null để đảm bảo chỉ người dùng đã xác thực mới có quyền truy cập dữ liệu.",
            "Đáp ứng yêu cầu 5 (Firestore Security Rules): Thay thế test-mode mở bằng luật bảo vệ tài nguyên, phân quyền đọc/ghi theo phiên đăng nhập."
        ),
        (
            4,
            "Ảnh 4: Màn hình Sign Up (Đăng ký tài khoản)",
            "Mở app trên điện thoại/máy ảo -> Bấm 'Create Account' -> Nhập thông tin form đăng ký (Full Name, Email, Password, Confirm Password). Chụp lại màn hình.",
            "Giao diện màn hình đăng ký tài khoản mới với các trường Name, Email, Password và xử lý kiểm tra dữ liệu hợp lệ.",
            "Đáp ứng yêu cầu 1 (Sign Up screen): Khởi tạo user trên Firebase Auth đồng thời tạo tài liệu user tương ứng trong collection users."
        ),
        (
            5,
            "Ảnh 5: Màn hình Login (Đăng nhập)",
            "Màn hình đăng nhập tài khoản với email, password và nút Sign In. Có xử lý kiểm tra rỗng và thông báo lỗi thân thiện.",
            "Giao diện màn hình đăng nhập tài khoản bằng Firebase Auth, có cơ chế ghi nhớ phiên đăng nhập qua onAuthStateChanged.",
            "Đáp ứng yêu cầu 1 (Login screen & State persistence): Tự động chuyển hướng vào ứng dụng khi đã đăng nhập thành công."
        ),
        (
            6,
            "Ảnh 6: Màn hình Home (Task List) sau khi đăng nhập",
            "Màn hình Home sau khi đăng nhập thành công, hiển thị danh sách công việc. Trên mỗi task card có gắn huy hiệu tên Team và tên người thực hiện (Assignee).",
            "Màn hình danh sách công việc được bảo vệ (sau đăng nhập), hiển thị các task kèm tên Team và tên Assignee.",
            "Đáp ứng yêu cầu 1 & 3: Bảo vệ màn hình chính sau xác thực, hiển thị trực quan thông tin phân quyền và gán việc."
        ),
        (
            7,
            "Ảnh 7: Màn hình Teams (Danh sách & Tạo / Tham gia Team)",
            "Chuyển sang tab Teams, hiển thị danh sách các team mà tài khoản đang tham gia kèm nút 'New' (tạo team cấp mã Join Code 6 ký tự) và nút 'Join' (tham gia bằng code).",
            "Màn hình Teams cho phép tạo team mới sinh mã Join Code và tham gia team bằng mã 6 ký tự.",
            "Đáp ứng yêu cầu 2 (Teams screen): Quản lý danh sách nhóm, hỗ trợ tạo nhóm mới và tham gia nhóm dễ dàng."
        ),
        (
            8,
            "Ảnh 8: Màn hình Team Detail (Chi tiết nhóm & Thành viên)",
            "Bấm vào 1 Team -> Mở Team Detail screen: Hiển thị tên nhóm, mã code, danh sách các thành viên (kèm role Owner/Member) và các task thuộc riêng nhóm đó.",
            "Màn hình chi tiết nhóm hiển thị danh sách thành viên với role (Owner/Member) và các task thuộc nhóm.",
            "Đáp ứng yêu cầu 2 (Team Detail screen): Hiển thị chi tiết thành viên và các nhiệm vụ gắn liền với nhóm."
        ),
        (
            9,
            "Ảnh 9: Màn hình Modal gán Task cho Team & Assignee",
            "Bấm nút (+) tạo task hoặc sửa task -> Hộp thoại TaskModal hiển thị picker chọn Team và dropdown chọn thành viên (Assignee) từ danh sách team đó.",
            "Form tạo/sửa task cho phép chọn Team và gán người thực hiện (Assignee) từ danh sách thành viên của team.",
            "Đáp ứng yêu cầu 3 (Tasks Linked to Teams & Assignees): Liên kết toàn diện giữa công việc, nhóm làm việc và người chịu trách nhiệm."
        ),
        (
            10,
            "Ảnh 10: Màn hình Real-Time Chat giữa 2 tài khoản (Alice & Bob)",
            "Từ màn hình Team Detail -> Bấm 'Open Real-Time Team Chat'. Chụp cuộc trò chuyện giữa 2 tài khoản: Phân biệt rõ tin nhắn gửi (bên phải, màu xanh) và tin nhắn nhận (bên trái, avatar người gửi, giờ gửi).",
            "Màn hình Chat thời gian thực với cuộc hội thoại giữa 2 user khác nhau, hiển thị rõ giờ gửi và phân biệt tin nhắn của mình (bên phải) và thành viên khác (bên trái kèm avatar).",
            "Đáp ứng yêu cầu 4 (Real-Time Chat): Chat subcollection teams/{teamId}/messages với Firestore onSnapshot cập nhật tức thì."
        ),
        (
            11,
            "Ảnh 11: Màn hình Profile và hành vi Logout",
            "Chuyển sang tab Profile -> Hiển thị thông tin người dùng đang đăng nhập -> Bấm nút đỏ 'Log Out of TaskFlow' -> Xác nhận ứng dụng đưa người dùng quay lại màn hình Login.",
            "Màn hình Profile hiển thị thông tin tài khoản hiện tại và xác nhận khi bấm Logout ứng dụng tự động quay về màn hình Login.",
            "Đáp ứng yêu cầu 1 (Logout & Auth protection): Đăng xuất an toàn, xóa trạng thái đăng nhập và bảo vệ điều hướng."
        )
    ]

    for num, title, guide, caption, req in screenshots_info:
        # Title of screenshot
        p_num = doc.add_paragraph()
        r_num = p_num.add_run(f"3.{num} {title}")
        r_num.font.name = "Arial"
        r_num.font.bold = True
        r_num.font.size = Pt(11.5)
        r_num.font.color.rgb = SECONDARY_COLOR

        # Placeholder box for image
        box = doc.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        box.autofit = False
        cell = box.rows[0].cells[0]
        cell.width = Inches(7.0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, 200, 200, 200, 200)

        p_box = cell.paragraphs[0]
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        r_box_icon = p_box.add_run("📸 [CHÈN ẢNH CHỤP MÀN HÌNH TẠI ĐÂY]\n\n")
        r_box_icon.font.name = "Arial"
        r_box_icon.font.size = Pt(12)
        r_box_icon.font.bold = True
        r_box_icon.font.color.rgb = PRIMARY_COLOR

        r_box_guide = p_box.add_run(f"👉 Thao tác tay cần chụp: {guide}\n(Paste ảnh chụp toàn màn hình rõ nét vào vị trí này)")
        r_box_guide.font.name = "Arial"
        r_box_guide.font.size = Pt(9.5)
        r_box_guide.font.color.rgb = RGBColor(71, 85, 105)

        # Official Caption below image
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap_label = p_cap.add_run(f"Caption: ")
        r_cap_label.font.name = "Arial"
        r_cap_label.font.size = Pt(10)
        r_cap_label.font.bold = True
        r_cap_label.font.color.rgb = PRIMARY_COLOR

        r_cap_txt = p_cap.add_run(f"{caption}")
        r_cap_txt.font.name = "Arial"
        r_cap_txt.font.size = Pt(10)
        r_cap_txt.font.italic = True
        r_cap_txt.font.color.rgb = SECONDARY_COLOR

        # Requirement tag
        p_req = doc.add_paragraph()
        r_req = p_req.add_run(f"✓ {req}")
        r_req.font.name = "Arial"
        r_req.font.size = Pt(9)
        r_req.font.color.rgb = ACCENT_COLOR

        doc.add_paragraph()

    # --- SECTION 4: CONCLUSION & VERIFICATION CHECKLIST ---
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("IV. BẢNG TỔNG HỢP KIỂM THỬ TÍNH NĂNG (FEATURE CHECKLIST)")
    h4_run.font.name = "Arial"
    h4_run.font.size = Pt(14)
    h4_run.font.bold = True
    h4_run.font.color.rgb = PRIMARY_COLOR

    chk_table = doc.add_table(rows=7, cols=3)
    chk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    chk_table.autofit = False

    # Header
    hdr = chk_table.rows[0]
    hdr.cells[0].width = Inches(2.2)
    hdr.cells[1].width = Inches(3.6)
    hdr.cells[2].width = Inches(1.2)
    for idx, text in enumerate(["Yêu cầu chức năng (Requirement)", "Mô tả triển khai trong Source Code", "Trạng thái"]):
        cell = hdr.cells[idx]
        set_cell_background(cell, "4F46E5")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    check_items = [
        ("1. Firebase Authentication", "Sign up, Login, Logout, onAuthStateChanged, User profile Firestore", "Đạt (100%)"),
        ("2. Teams Management", "Tạo nhóm, mã code 6 ký tự, tham gia nhóm bằng code, Team Detail", "Đạt (100%)"),
        ("3. Task Linking & Assignee", "Gán Task cho Team, chọn Assignee từ thành viên, hiển thị trên card", "Đạt (100%)"),
        ("4. Real-Time Team Chat", "Subcollection teams/{teamId}/messages, onSnapshot listener, bubble UI", "Đạt (100%)"),
        ("5. Firestore Security Rules", "Yêu cầu request.auth != null bảo vệ toàn diện dữ liệu", "Đạt (100%)"),
        ("6. Bonus Features", "ERD schema trong README, responsive chat bubbles, UI polishing", "Đạt (100%)")
    ]

    for r_idx, (r_name, r_desc, r_st) in enumerate(check_items):
        row = chk_table.rows[r_idx+1]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(3.6)
        row.cells[2].width = Inches(1.2)
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"

        for c_idx, val in enumerate([r_name, r_desc, r_st]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = SECONDARY_COLOR
            elif c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = ACCENT_COLOR

    doc.add_paragraph()
    doc.add_paragraph()

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig1 = p_sig.add_run("Sinh viên thực hiện\n(Ký và ghi rõ họ tên)\n\n\n\n")
    r_sig1.font.name = "Arial"
    r_sig1.font.size = Pt(10)
    r_sig1.font.italic = True
    r_sig1.font.color.rgb = MUTED_COLOR

    r_sig2 = p_sig.add_run("[Họ và tên sinh viên]")
    r_sig2.font.name = "Arial"
    r_sig2.font.size = Pt(11)
    r_sig2.font.bold = True
    r_sig2.font.color.rgb = SECONDARY_COLOR

    output_path = r"d:\MMA301\StudentID_Exam2.docx"
    doc.save(output_path)
    print(f"Report document successfully generated at: {output_path}")

if __name__ == "__main__":
    create_report()
